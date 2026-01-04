from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ["Program Name", "Starting Date"]

def autosize(ws):
    for col in range(1, ws.max_column + 1):
        letter = get_column_letter(col)
        max_len = 0
        for cell in ws[letter]:
            if cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[letter].width = min(max_len + 2, 60)

def write_excel(output_path: str, data_by_sheet: dict[str, list[dict]], logger):
    Path(output_path).parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()

    # remove default
    default_ws = wb.active
    wb.remove(default_ws)

    for sheet_name, rows in data_by_sheet.items():
        safe_name = sheet_name[:31]  # Excel limit
        ws = wb.create_sheet(title=safe_name)

        # headers
        for i, h in enumerate(HEADERS, start=1):
            cell = ws.cell(row=1, column=i, value=h)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        # data
        r = 2
        for row in rows:
            ws.cell(row=r, column=1, value=row.get("Program Name", ""))
            ws.cell(row=r, column=2, value=row.get("Starting Date", ""))
            r += 1

        ws.freeze_panes = "A2"
        autosize(ws)

    wb.save(output_path)
    logger.info(f"Excel saved: {output_path}")
